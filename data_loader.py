"""
Data Loader: reads all Excel files and constructs Agent objects.

Loading sequence:
  1. Phase_4B (condition scores + categories) → internal/external conditions
  2. Phase_4D (outcome scores + categories) → outcomes
  3. Phase_1A (case equations) → production rules
  4. Relational_Density → relational profiles
  5. Phase_3A/3B (cross-impact matrices) → validation reference data
  6. Condition/Outcome mappings → reference lookups
  7. Sovereignty profiles computed from outcome categories
"""

import re
import openpyxl
from collections import defaultdict
import json
from config import (
    FILES, CASE_NAME_MAP, CANONICAL_CASES,
    CONDITION_TYPE_MAP, CONDITION_COL_NAMES, CONDITION_CATEGORIES,
    CONDITION_MECHANISM_MAP, INTERNAL_CONDITION_TYPES, EXTERNAL_CONDITION_TYPES,
    OUTCOME_COL_NAMES, OUTCOME_CATEGORIES, OUTCOME_SOVEREIGNTY_MAP,
    SOVEREIGNTY_DIMENSIONS, ORG_TYPE_MAP, CONDITION_SOVEREIGNTY_TERRAIN,
    RECURSIVE_PATTERNS, MECHANISM_RECURSIVE_MAP, RECURSIVE_PATTERN_REFINEMENT,
    canonicalize_case_name, canonicalize_condition_type, canonicalize_org_type,
)
from agent import Agent, ProductionRuleRow, RelationalProfile


class DataLoader:
    """Loads and integrates all data sources into Agent objects."""

    def __init__(self):
        self.agents = {}              # {canonical_name: Agent}
        self.cross_impact_quant = {}  # Quantitative cross-impact matrix
        self.cross_impact_qual = {}   # Qualitative cross-impact matrix
        self.pattern_summary = {}     # Mechanism type patterns
        self.condition_ref = {}       # Condition category reference
        self.outcome_ref = {}         # Outcome category reference

    def load_all(self) -> dict:
        """Load all data and return dict of agents."""
        print("Loading Phase_4B (condition positioning)...")
        self._load_phase_4b()
        print(f"  → {len(self.agents)} agents created with condition scores")

        print("Loading Phase_4D (outcome positioning)...")
        self._load_phase_4d()
        print(f"  → Outcome scores loaded")

        print("Loading Phase_1A (case equations → production rules)...")
        self._load_phase_1a()
        total_rules = sum(len(a.production_rule) for a in self.agents.values())
        print(f"  → {total_rules} production rule rows loaded")

        print("Loading Relational_Density...")
        self._load_relational_density()
        print(f"  → Relational profiles loaded")

        print("Loading cross-impact matrices...")
        self._load_cross_impact()
        print(f"  → Quantitative: {len(self.cross_impact_quant)} outcome rows")
        print(f"  → Qualitative: {len(self.cross_impact_qual)} outcome rows")

        print("Loading reference mappings...")
        self._load_reference_mappings()
        print(f"  → {len(self.condition_ref)} condition categories")
        print(f"  → {len(self.outcome_ref)} outcome categories")

        print("Loading Appendix 3 enrichment...")
        self._load_appendix3_enrichment()
        print(f"  → Case equation narratives loaded")

        print("Computing derived characteristics...")
        self._compute_derived()
        print(f"  → Condition diversity, outcome breadth, sovereignty profiles")

        print("Assigning recursive D-A-S patterns...")
        self._assign_recursive_patterns()
        print(f"  → 17 recursive patterns assigned at agent level")

        print("Validating...")
        self._validate()

        return self.agents

    # ── Phase_4B: Condition Positioning ───────────────────────────────

    def _load_phase_4b(self):
        """Load condition scores and categories from Phase_4B."""
        wb = openpyxl.load_workbook(FILES["phase_4b"], read_only=True)

        # --- Scored Positioning ---
        ws = wb["Scored Positioning"]
        rows = list(ws.iter_rows(values_only=True))
        headers = rows[0]

        # Map column indices to condition type codes
        col_map = {}
        for i, h in enumerate(headers):
            if h and h in CONDITION_COL_NAMES:
                col_map[i] = CONDITION_COL_NAMES[h]

        for row in rows[1:]:
            raw_name = str(row[0]).strip()
            canonical = canonicalize_case_name(raw_name)
            org_type = canonicalize_org_type(str(row[1]).strip())

            agent = Agent(name=canonical, org_type=org_type)

            for col_idx, ctype in col_map.items():
                val = row[col_idx]
                score = self._parse_score(val)
                if ctype in INTERNAL_CONDITION_TYPES:
                    agent.internal_condition_scores[ctype] = score
                else:
                    agent.external_condition_scores[ctype] = score

            self.agents[canonical] = agent

        # --- Categorical Matrix ---
        ws2 = wb["Categorical Matrix"]
        rows2 = list(ws2.iter_rows(values_only=True))
        headers2 = rows2[0]

        cat_col_map = {}
        for i, h in enumerate(headers2):
            if h and h in CONDITION_COL_NAMES:
                cat_col_map[i] = CONDITION_COL_NAMES[h]

        for row in rows2[1:]:
            raw_name = str(row[0]).strip()
            canonical = canonicalize_case_name(raw_name)
            agent = self.agents[canonical]

            for col_idx, ctype in cat_col_map.items():
                val = row[col_idx]
                codes = self._parse_category_codes(val)
                if codes:
                    agent.condition_categories[ctype] = codes

        wb.close()

    # ── Phase_4D: Outcome Positioning ─────────────────────────────────

    def _load_phase_4d(self):
        """Load outcome scores and categories from Phase_4D."""
        wb = openpyxl.load_workbook(FILES["phase_4d"], read_only=True)

        # --- Scored Positioning ---
        ws = wb["Scored Positioning"]
        rows = list(ws.iter_rows(values_only=True))
        headers = rows[0]

        col_map = {}
        for i, h in enumerate(headers):
            if h and h in OUTCOME_COL_NAMES:
                col_map[i] = OUTCOME_COL_NAMES[h]

        # Find Total and Groups columns
        total_col = None
        groups_col = None
        for i, h in enumerate(headers):
            if h == "Total":
                total_col = i
            if h == "Groups":
                groups_col = i

        for row in rows[1:]:
            raw_name = str(row[0]).strip()
            canonical = canonicalize_case_name(raw_name)
            agent = self.agents[canonical]

            for col_idx, ogroup in col_map.items():
                val = row[col_idx]
                agent.outcome_scores[ogroup] = self._parse_score(val)

            if total_col is not None:
                agent.outcome_total = self._parse_score(row[total_col]) or 0.0
            if groups_col is not None:
                val = row[groups_col]
                agent.outcome_breadth = int(val) if val and val != "N/A" else 0

        # --- Categorical Matrix ---
        ws2 = wb["Categorical Matrix"]
        rows2 = list(ws2.iter_rows(values_only=True))
        headers2 = rows2[0]

        cat_col_map = {}
        for i, h in enumerate(headers2):
            if h and h in OUTCOME_COL_NAMES:
                cat_col_map[i] = OUTCOME_COL_NAMES[h]

        for row in rows2[1:]:
            raw_name = str(row[0]).strip()
            canonical = canonicalize_case_name(raw_name)
            agent = self.agents[canonical]

            for col_idx, ogroup in cat_col_map.items():
                val = row[col_idx]
                codes = self._parse_category_codes(val)
                if codes:
                    agent.outcome_categories[ogroup] = codes

        wb.close()

    # ── Phase_1A: Case Equations → Production Rules ───────────────────

    def _load_phase_1a(self):
        """Load case equations from Phase_1A and build production rules."""
        wb = openpyxl.load_workbook(FILES["phase_1a"], read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        headers = rows[0]
        # Headers: Case Name, Description, Condition Type, Condition Name,
        #          Challenge Addressed, Outcome Group, Specific Outcome Produced,
        #          Organization Type, Economic Sector, Location, Geographic Scope, Maturity

        for row in rows[1:]:
            raw_name = str(row[0]).strip()
            canonical = canonicalize_case_name(raw_name)
            agent = self.agents[canonical]

            # Set characteristics (may be set multiple times, same values)
            if row[8]:
                agent.economic_sector = str(row[8]).strip()
            if row[9]:
                agent.location = str(row[9]).strip()
            if row[10]:
                agent.geographic_scope = str(row[10]).strip()
            if row[11]:
                agent.maturity = str(row[11]).strip()

            # Build production rule row
            cond_type_raw = str(row[2]).strip() if row[2] else ""
            cond_type_code = canonicalize_condition_type(cond_type_raw) if cond_type_raw else ""
            cond_name = str(row[3]).strip() if row[3] else ""
            challenge = str(row[4]).strip() if row[4] else ""
            outcome_group_raw = str(row[5]).strip() if row[5] else ""
            outcome_name = str(row[6]).strip() if row[6] else ""
            description = str(row[1]).strip() if row[1] else ""

            # Resolve outcome group to code
            outcome_group_code = OUTCOME_COL_NAMES.get(outcome_group_raw, "")

            # Determine mechanism type from condition type
            mechanism = CONDITION_MECHANISM_MAP.get(cond_type_code, "enabling")

            # Resolve condition codes from agent's categorical matrix
            cond_codes = agent.condition_categories.get(cond_type_code, [])

            # Resolve outcome codes from agent's categorical matrix
            outcome_codes = agent.outcome_categories.get(outcome_group_code, [])

            pr_row = ProductionRuleRow(
                condition_type=cond_type_code,
                condition_name=cond_name,
                condition_codes=list(cond_codes),
                mechanism_type=mechanism,
                outcome_group=outcome_group_code,
                outcome_name=outcome_name,
                outcome_codes=list(outcome_codes),
                description=description,
                challenge=challenge,
            )
            agent.production_rule.append(pr_row)

        wb.close()

    # ── Relational Density ────────────────────────────────────────────

    def _load_relational_density(self):
        """Load relational density scoring for all 43 cases."""
        wb = openpyxl.load_workbook(FILES["relational_density"], read_only=True)
        ws = wb["Relational Density Scoring"]
        rows = list(ws.iter_rows(values_only=True))

        for row in rows[1:]:
            if not row[0]:
                continue
            raw_name = str(row[0]).strip()
            canonical = canonicalize_case_name(raw_name)
            agent = self.agents.get(canonical)
            if not agent:
                print(f"  WARNING: Relational density case not found: {raw_name} → {canonical}")
                continue

            # Parse partner types engaged (text like "Indigenous Gov (1), Academic (3), ...")
            partner_types = self._parse_partner_types(row[2])

            # Parse PT categories
            pt_cats = self._parse_category_codes(row[5])

            profile = RelationalProfile(
                partner_types_engaged=partner_types,
                relationship_types=pt_cats,
                partner_diversity=int(row[3]) if row[3] else 0,
                partner_diversity_norm=float(row[4]) if row[4] else 0.0,
                relationship_type_range=int(row[6]) if row[6] else 0,
                relationship_type_range_norm=float(row[7]) if row[7] else 0.0,
                network_reach_description=str(row[8]) if row[8] else "",
                network_reach=int(row[9]) if row[9] else 0,
                network_reach_norm=float(row[10]) if row[10] else 0.0,
                inter_case_description=str(row[11]) if row[11] else "",
                inter_case_connectivity=int(row[12]) if row[12] else 0,
                inter_case_connectivity_norm=float(row[13]) if row[13] else 0.0,
                composite_score=float(row[14]) if row[14] else 0.0,
            )

            # Set direct partners from inter-case connections
            from config import DIRECT_PARTNERSHIPS
            for a, b in DIRECT_PARTNERSHIPS:
                if canonical == a:
                    profile.direct_partners.append(b)
                elif canonical == b:
                    profile.direct_partners.append(a)

            # Set latent clusters
            from config import LATENT_CLUSTERS
            for cluster_name, members in LATENT_CLUSTERS.items():
                if canonical in members:
                    profile.latent_clusters.append(cluster_name)

            agent.relational_profile = profile
            agent.relational_density = profile.composite_score

    # ── Cross-Impact Matrices ─────────────────────────────────────────

    def _load_cross_impact(self):
        """Load quantitative and qualitative cross-impact matrices."""
        # --- Quantitative (Phase_3A) ---
        wb = openpyxl.load_workbook(FILES["phase_3a_quant"], read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        headers = rows[0]

        # Extract condition codes from headers
        cond_codes = []
        for h in headers[1:]:
            if h and ":" in str(h):
                code = str(h).split(":")[0].strip()
                if code.startswith(("TP", "RS", "GV", "TT", "PT", "PU", "OI", "PO", "RG", "SC", "PR")):
                    cond_codes.append(code)
                else:
                    break
            elif h and str(h).startswith(("ROW_TOTAL", "SUB_")):
                break
            else:
                cond_codes.append(None)

        for row in rows[1:]:
            outcome_raw = str(row[0]).strip() if row[0] else ""
            if not outcome_raw or outcome_raw.startswith("SUBTOTAL") or outcome_raw.startswith("GRAND"):
                continue
            outcome_code = outcome_raw.split(":")[0].strip()
            self.cross_impact_quant[outcome_code] = {}
            for i, ccode in enumerate(cond_codes):
                if ccode and i + 1 < len(row):
                    val = row[i + 1]
                    self.cross_impact_quant[outcome_code][ccode] = int(val) if val else 0

        wb.close()

        # --- Qualitative (Phase_3B) ---
        wb = openpyxl.load_workbook(FILES["phase_3b_qual"], read_only=True)
        ws = wb["Qualitative Matrix"]
        rows = list(ws.iter_rows(values_only=True))
        headers = rows[0]

        qual_cond_codes = []
        for h in headers[1:]:
            if h and ":" in str(h):
                code = str(h).split(":")[0].strip()
                qual_cond_codes.append(code)
            else:
                qual_cond_codes.append(None)

        for row in rows[1:]:
            outcome_raw = str(row[0]).strip() if row[0] else ""
            if not outcome_raw:
                continue
            outcome_code = outcome_raw.split(":")[0].strip()
            self.cross_impact_qual[outcome_code] = {}
            for i, ccode in enumerate(qual_cond_codes):
                if ccode and i + 1 < len(row):
                    val = row[i + 1]
                    descriptor = str(val).strip() if val else "—"
                    if descriptor != "—" and descriptor != "None":
                        self.cross_impact_qual[outcome_code][ccode] = descriptor

        # Pattern Summary
        ws2 = wb["Pattern Summary"]
        rows2 = list(ws2.iter_rows(values_only=True))
        for row in rows2[1:]:
            if row[0] and row[2]:
                ctype = str(row[0]).strip()
                self.pattern_summary[ctype] = {
                    "primary_actions": str(row[1]).strip() if row[1] else "",
                    "dominant_mechanism": str(row[2]).strip() if row[2] else "",
                    "key_descriptors": str(row[3]).strip() if row[3] else "",
                }

        wb.close()

    # ── Reference Mappings ────────────────────────────────────────────

    def _load_reference_mappings(self):
        """Load condition and outcome category reference data."""
        # Condition Mapping
        wb = openpyxl.load_workbook(FILES["condition_mapping"], read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        for row in rows[1:]:
            if row[0]:
                code = str(row[0]).strip()
                self.condition_ref[code] = {
                    "type": str(row[1]).strip() if row[1] else "",
                    "category": str(row[2]).strip() if row[2] else "",
                    "cases": str(row[3]).strip() if row[3] else "",
                    "count": int(row[4]) if row[4] else 0,
                }
        wb.close()

        # Outcome Mapping
        wb = openpyxl.load_workbook(FILES["outcome_mapping"], read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        for row in rows[1:]:
            if row[0]:
                code = str(row[0]).strip()
                self.outcome_ref[code] = {
                    "group": str(row[1]).strip() if row[1] else "",
                    "category": str(row[2]).strip() if row[2] else "",
                    "cases": str(row[3]).strip() if row[3] else "",
                    "count": int(row[4]) if row[4] else 0,
                }
        wb.close()

    # ── Compute Derived Characteristics ───────────────────────────────

    def _load_appendix3_enrichment(self):
        """Load condition interaction, synthesis, and compact equation from Appendix 3."""
        import os
        json_path = os.path.join(os.path.dirname(__file__), 'appendix3_enrichment.json')
        if not os.path.exists(json_path):
            print("    (appendix3_enrichment.json not found, skipping)")
            return
        with open(json_path, 'r') as f:
            enrichment = json.load(f)
        loaded = 0
        for name, data in enrichment.items():
            if name in self.agents:
                agent = self.agents[name]
                agent.condition_interactions = data.get('condition_interactions', '')
                agent.synthesis = data.get('synthesis', '')
                agent.compact_equation = data.get('compact_equation', '')
                loaded += 1
        if loaded != len(self.agents):
            missing = set(self.agents.keys()) - set(enrichment.keys())
            if missing:
                print(f"    Warning: {len(missing)} agents without enrichment: {missing}")

    def _assign_recursive_patterns(self):
        """
        Assign 17 recursive D-A-S patterns at agent level (BINARY).

        Each production rule row has a mechanism type. The mechanism constrains
        which patterns are reachable. Condition type + outcome group context
        selects specific patterns from the reachable set. A case either
        participates in a pattern or it does not — no counting.
        """
        for agent in self.agents.values():
            pattern_set = set()
            for row in agent.production_rule:
                mechanism = row.mechanism_type.split("+")[0].strip()
                # Try refined mapping first (constrained to mechanism's reachable set)
                key = (row.condition_type, row.outcome_group)
                refined = RECURSIVE_PATTERN_REFINEMENT.get(key)
                if refined:
                    pattern_set.update(refined)
                else:
                    # Fall back to mechanism-level mapping (all reachable patterns)
                    candidates = MECHANISM_RECURSIVE_MAP.get(mechanism, [])
                    pattern_set.update(candidates)

            agent.recursive_patterns = sorted(pattern_set)  # list of pattern IDs
            agent.dominant_patterns = agent.recursive_patterns  # all patterns (binary, no ranking)
            agent.dominant_pattern_names = [
                RECURSIVE_PATTERNS.get(p, f"#{p}") for p in agent.recursive_patterns
            ]

    def _compute_derived(self):
        """Compute condition diversity, sovereignty profiles, etc."""
        for agent in self.agents.values():
            # Condition diversity: count of condition types with non-None scores
            agent.condition_diversity = len(agent.active_condition_types)

            # Outcome breadth: already loaded from Phase_4D 'Groups' column,
            # but verify/compute if missing
            if agent.outcome_breadth == 0:
                agent.outcome_breadth = len(agent.active_outcome_groups)

            # Sovereignty profile: computed from outcome categories (Bridge D)
            agent.sovereignty_profile = self._compute_sovereignty_profile(agent)

            # Sovereignty terrain: computed from condition types (Bridge B)
            agent.sovereignty_terrain = self._compute_sovereignty_terrain(agent)

    def _compute_sovereignty_profile(self, agent: Agent) -> dict:
        """
        Compute sovereignty capacity profile from outcome categories.
        Each outcome category maps to sovereignty dimensions via
        OUTCOME_SOVEREIGNTY_MAP. The profile aggregates contributions.
        """
        profile = {dim: 0.0 for dim in SOVEREIGNTY_DIMENSIONS}
        contribution_count = {dim: 0 for dim in SOVEREIGNTY_DIMENSIONS}

        for ogroup, codes in agent.outcome_categories.items():
            group_score = agent.outcome_scores.get(ogroup)
            if group_score is None:
                continue
            for code in codes:
                dims = OUTCOME_SOVEREIGNTY_MAP.get(code, [])
                for dim in dims:
                    profile[dim] += group_score
                    contribution_count[dim] += 1

        # Normalize: average contribution per dimension
        for dim in SOVEREIGNTY_DIMENSIONS:
            if contribution_count[dim] > 0:
                profile[dim] = round(profile[dim] / contribution_count[dim], 3)

        return profile

    def _compute_sovereignty_terrain(self, agent: Agent) -> dict:
        """
        Compute sovereignty terrain profile from condition types (Bridge B).
        Maps condition types to the sovereignty dimensions they OPERATE ON,
        weighted by condition scores. Distinct from sovereignty_profile
        (Bridge D) which maps outcomes to constructed sovereignty capacity.
        """
        terrain = {dim: 0.0 for dim in SOVEREIGNTY_DIMENSIONS}
        contribution_count = {dim: 0 for dim in SOVEREIGNTY_DIMENSIONS}

        for ctype, score in agent.all_condition_scores.items():
            if score is None:
                continue
            dims = CONDITION_SOVEREIGNTY_TERRAIN.get(ctype, [])
            for dim in dims:
                terrain[dim] += score
                contribution_count[dim] += 1

        # Normalize: average contribution per dimension
        for dim in SOVEREIGNTY_DIMENSIONS:
            if contribution_count[dim] > 0:
                terrain[dim] = round(terrain[dim] / contribution_count[dim], 3)

        return terrain

    # ── Validation ────────────────────────────────────────────────────

    def _validate(self):
        """Run validation checks on loaded data."""
        errors = []
        warnings = []

        if len(self.agents) != 43:
            errors.append(f"Expected 43 agents, got {len(self.agents)}")

        for name, agent in self.agents.items():
            # Check condition scores exist
            if not agent.internal_condition_scores and not agent.external_condition_scores:
                errors.append(f"{name}: No condition scores loaded")

            # Check at least some outcomes
            if not agent.outcome_scores:
                errors.append(f"{name}: No outcome scores loaded")

            # Check production rule exists
            if not agent.production_rule:
                warnings.append(f"{name}: No production rule rows")

            # Check relational profile
            if agent.relational_profile is None:
                warnings.append(f"{name}: No relational profile")

            # Verify condition diversity matches
            active = len(agent.active_condition_types)
            if agent.condition_diversity != active:
                warnings.append(
                    f"{name}: condition_diversity={agent.condition_diversity} "
                    f"but active types={active}"
                )

        if errors:
            print(f"\n  ❌ ERRORS ({len(errors)}):")
            for e in errors:
                print(f"    {e}")
        if warnings:
            print(f"\n  ⚠ WARNINGS ({len(warnings)}):")
            for w in warnings[:10]:
                print(f"    {w}")
            if len(warnings) > 10:
                print(f"    ... and {len(warnings) - 10} more")

        if not errors:
            print(f"  ✓ All {len(self.agents)} agents validated successfully")

    # ── Parsing Helpers ───────────────────────────────────────────────

    @staticmethod
    def _parse_score(val) -> float | None:
        """Parse a score value, returning None for N/A."""
        if val is None or str(val).strip().upper() in ("N/A", "NA", "—", "-", ""):
            return None
        try:
            return float(val)
        except (ValueError, TypeError):
            return None

    @staticmethod
    def _parse_category_codes(val) -> list:
        """Parse category codes from a cell value like 'DP1, DP8' or '—'."""
        if val is None:
            return []
        val_str = str(val).strip()
        if val_str in ("—", "-", "N/A", "", "None"):
            return []
        # Split on comma and clean
        codes = []
        for part in val_str.split(","):
            code = part.strip()
            # Validate it looks like a category code
            if re.match(r"^[A-Z]{2}\d+$", code):
                codes.append(code)
        return codes

    @staticmethod
    def _parse_partner_types(val) -> list:
        """Parse partner type numbers from text like 'Indigenous Gov (1), Academic (3)'."""
        if not val:
            return []
        numbers = re.findall(r"\((\d+)\)", str(val))
        return [int(n) for n in numbers]


# ── Convenience Function ──────────────────────────────────────────────

def load_model_data():
    """Load all data and return (agents_dict, loader)."""
    loader = DataLoader()
    agents = loader.load_all()
    return agents, loader
